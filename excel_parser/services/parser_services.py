import os
import openpyxl

from excel_parser.models import ExcelDocument, DocumentProcessingStatuses


class Document():

    def __init__(self, document_id: int):
        self.document = ExcelDocument.objects.get(id=document_id)
        self.wb = openpyxl.load_workbook(os.path.join(self.document.document.path))
        self.document.processing_status = DocumentProcessingStatuses.IN_PROCESS
        self.document.save()

    def _get_actual_sheet_name(self) -> dict:
        '''
        Находит лист в документе, где присутствуеют нужные столбцы
        '''
        actual_sheet_data = dict()
        for sheet_name in self.wb.sheetnames:
            if not self.wb[sheet_name]['A1'].value:
                continue
            else:
                for cell in self.wb[sheet_name]['1']:
                    if cell.value == 'after' or cell.value == 'before':
                        actual_sheet_data[cell.value] = cell.column_letter
                        actual_sheet_data['sheet_name'] = sheet_name
                    else:
                        continue

        return actual_sheet_data

    def get_result(self) -> None:
        '''
        Получает значение X и пишет его в БД
        '''
        actual_sheet_data = self._get_actual_sheet_name()
        if not actual_sheet_data:
            self.document.save_result(result='Incorrect file or empty values')
            self.wb.close()

        after_list = list()
        before_list = list()
        for coll in self.wb[actual_sheet_data['sheet_name']][actual_sheet_data['after']]:
            after_list.append(coll.value)

        for coll in self.wb[actual_sheet_data['sheet_name']][actual_sheet_data['before']]:
            before_list.append(coll.value)

        if (after_list[1] - before_list[1]) > 0:
            self.document.save_result(result=f'added: {after_list[1] - before_list[1]}')
        else:
            self.document.save_result(result=f'removed: {before_list[1] - after_list[1]}')

        self.wb.close()
